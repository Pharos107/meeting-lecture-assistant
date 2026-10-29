"""Xuất Markdown, DOCX, PDF, SRT, và VTT."""

import io
import re
from pathlib import Path
from urllib.parse import quote

from docx import Document
from fpdf import FPDF
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.routes_jobs import current_segments, load_aliases, speaker_name
from app.models import ActionItem, Highlight, Job, LlmRun, Segment

_FONT = Path(__file__).resolve().parents[2] / "assets" / "fonts" / "NotoSans-Regular.ttf"


def choose_text(segment: Segment, mode: str) -> str:
    if mode == "raw":
        return segment.text_raw
    if mode == "cleaned":
        if segment.text_cleaned is not None:
            return segment.text_cleaned
        return segment.text_raw
    if mode == "edited":
        if segment.text_edited is not None:
            return segment.text_edited
        if segment.text_cleaned is not None:
            return segment.text_cleaned
        return segment.text_raw
    if segment.text_edited is not None:
        return segment.text_edited
    if segment.text_cleaned is not None:
        return segment.text_cleaned
    return segment.text_raw


def clock(ms: int | None) -> str:
    total = max(int(ms or 0), 0) // 1000
    minutes, seconds = divmod(total, 60)
    return f"{minutes:02d}:{seconds:02d}"


def srt_stamp(ms: int) -> str:
    value = max(int(ms), 0)
    hours, rem = divmod(value, 3_600_000)
    minutes, rem = divmod(rem, 60_000)
    seconds, millis = divmod(rem, 1000)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d},{millis:03d}"


def vtt_stamp(ms: int) -> str:
    return srt_stamp(ms).replace(",", ".")


def safe_title(title: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\r\n]+', " ", title).strip()
    return cleaned or "transcript"


def content_disposition(title: str, extension: str) -> str:
    filename = f"{safe_title(title)}.{extension}"
    return f"attachment; filename*=UTF-8''{quote(filename)}"


def _cue_text(segment: Segment, aliases: dict, mode: str, include_speakers: bool) -> str:
    text = choose_text(segment, mode)
    if not include_speakers:
        return text
    name = speaker_name(segment.speaker_label, aliases)
    if not name:
        return text
    return f"{name}: {text}"


def subtitle_cues(segments: list[Segment], aliases: dict, mode: str, include_speakers: bool, stamp) -> str:
    blocks = []
    for index, segment in enumerate(segments, start=1):
        timing = f"{stamp(segment.start_ms)} --> {stamp(segment.end_ms)}"
        text = _cue_text(segment, aliases, mode, include_speakers)
        if stamp is srt_stamp:
            blocks.append(f"{index}\n{timing}\n{text}\n")
        else:
            blocks.append(f"{timing}\n{text}\n")
    return "\n".join(blocks)


def render_srt(segments, aliases, mode: str, include_speakers: bool) -> str:
    return subtitle_cues(segments, aliases, mode, include_speakers, srt_stamp)


def render_vtt(segments, aliases, mode: str, include_speakers: bool) -> str:
    body = subtitle_cues(segments, aliases, mode, include_speakers, vtt_stamp)
    return f"WEBVTT\n\n{body}"


def _lines_of(items: list, formatter) -> list[str]:
    return [formatter(item) for item in items if formatter(item)]


def document_sections(db: Session, job: Job, segments: list[Segment], aliases: dict, mode: str, include_speakers: bool) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = []
    run = db.scalar(
        select(LlmRun)
        .where(LlmRun.job_id == job.id, LlmRun.task == "summarize", LlmRun.status == "done")
        .order_by(LlmRun.created_at.desc())
    )
    result = run.result if run is not None and isinstance(run.result, dict) else None
    if run is not None and run.template == "meeting" and result is not None:
        sections.append(("Mục tiêu", [str(result.get("objective") or "")]))
        sections.append(("Vấn đề tồn đọng", _lines_of(result.get("open_issues") or [], _stamp_text)))
        sections.append(("Quyết định", _lines_of(result.get("decisions") or [], _stamp_text)))
        sections.append(("Việc do biên bản đề xuất", _lines_of(result.get("actions") or [], _meeting_action)))
    elif run is not None and run.template == "lecture" and result is not None:
        sections.append(("Tóm tắt", [str(result.get("overview") or "")]))
        sections.append(("Khái niệm", _lines_of(result.get("concepts") or [], _concept)))
        sections.append(("Ví dụ", _lines_of(result.get("examples") or [], _stamp_text)))
        sections.append(("Dặn dò ôn", _lines_of(result.get("exam_notes") or [], _stamp_text)))

    actions = list(
        db.scalars(
            select(ActionItem).where(ActionItem.job_id == job.id, ActionItem.source == "manual")
        ).all()
    )
    actions.sort(key=lambda row: ((row.start_ms or 0), str(row.id)))
    if actions:
        lines = []
        for action in actions:
            extra = " — ".join(part for part in (action.owner, action.deadline) if part)
            line = f"{clock(action.start_ms)} {action.task}"
            if extra:
                line = f"{line} — {extra}"
            lines.append(line)
        sections.append(("Việc cần làm", lines))

    highlights = list(db.scalars(select(Highlight).where(Highlight.job_id == job.id)).all())
    highlights.sort(key=lambda row: (row.start_ms, str(row.id)))
    if highlights:
        sections.append(("Đánh dấu", [f"{clock(row.start_ms)} {row.quote}" for row in highlights]))

    transcript = [
        f"[{clock(segment.start_ms)}] {_cue_text(segment, aliases, mode, include_speakers)}"
        for segment in segments
    ]
    sections.append(("Transcript", transcript))
    return sections


def _stamp_text(item: dict) -> str:
    if not isinstance(item, dict):
        return ""
    return f"{item.get('timestamp') or ''} {item.get('text') or ''}".strip()


def _meeting_action(item: dict) -> str:
    if not isinstance(item, dict):
        return ""
    extra = " — ".join(part for part in (item.get("owner") or "", item.get("deadline") or "") if part)
    line = f"{item.get('timestamp') or ''} {item.get('task') or ''}".strip()
    if extra:
        return f"{line} — {extra}"
    return line


def _concept(item: dict) -> str:
    if not isinstance(item, dict):
        return ""
    return f"{item.get('timestamp') or ''} {item.get('term') or ''}: {item.get('explanation') or ''}".strip()


def render_markdown(sections: list[tuple[str, list[str]]]) -> str:
    parts = []
    for heading, lines in sections:
        parts.append(f"# {heading}")
        parts.extend(line for line in lines if line)
        parts.append("")
    return "\n".join(parts).rstrip() + "\n"


def render_docx(sections: list[tuple[str, list[str]]]) -> bytes:
    document = Document()
    for heading, lines in sections:
        document.add_heading(heading, level=1)
        for line in lines:
            if line:
                document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_pdf(sections: list[tuple[str, list[str]]]) -> bytes:
    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.add_font("Noto", "", str(_FONT))
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    for heading, lines in sections:
        pdf.set_font("Noto", size=16)
        pdf.multi_cell(0, 10, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Noto", size=12)
        for line in lines:
            if line:
                pdf.multi_cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)
    return bytes(pdf.output())


def export_job(db: Session, job: Job, fmt: str, mode: str, include_speakers: bool) -> tuple[bytes, str, str]:
    _pass_no, segments = current_segments(db, job)
    aliases = load_aliases(db, job)
    if fmt == "srt":
        body = render_srt(segments, aliases, mode, include_speakers).encode("utf-8")
        return body, "application/x-subrip; charset=utf-8", "srt"
    if fmt == "vtt":
        body = render_vtt(segments, aliases, mode, include_speakers).encode("utf-8")
        return body, "text/vtt; charset=utf-8", "vtt"
    sections = document_sections(db, job, segments, aliases, mode, include_speakers)
    if fmt == "md":
        return render_markdown(sections).encode("utf-8"), "text/markdown; charset=utf-8", "md"
    if fmt == "docx":
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        return render_docx(sections), media, "docx"
    return render_pdf(sections), "application/pdf", "pdf"
